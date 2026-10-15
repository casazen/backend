"""Esporta il modello in Excel con formule vive.

Il foglio "Ipotesi" contiene tutti i parametri (celle gialle). Il foglio "Mensile" ricalcola i 36 mesi
dallo scenario scelto in Ipotesi!C4. Il foglio "Sintesi" somma per anno.
La logica replica `modello_economico.run`; `verifica_xlsx.py` confronta i due risultati.
"""

from __future__ import annotations

from dataclasses import fields

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from modello_economico import MONTHS, FixedCosts, Scenario, UnitEconomics, month_label

INPUT = PatternFill("solid", fgColor="FFF2CC")
HEAD = Font(bold=True)

UE_LABELS = {
    "list_all_in_small": "Listino tutto compreso 10-49 unità (€/unità)",
    "list_all_in_mid": "Listino tutto compreso 50-149 unità (€/unità)",
    "list_all_in_large": "Listino tutto compreso 150-399 unità (€/unità)",
    "verifications_per_unit": "Verifiche d'identità per unità al mese",
    "fp_fee_per_unit": "Canone Founding Partner (€/unità)",
    "nextpax_min": "NextPax minimo mensile (€)",
    "nextpax_included_units": "Unità incluse nel minimo NextPax",
    "nextpax_per_unit": "NextPax oltre il minimo (€/unità)",
    "services_per_unit": "SDI, AI, video (€/unità)",
    "verification_cost": "Costo di una verifica (€)",
    "b2c_arpu": "Canone B2C (€/host) - IPOTESI",
    "b2c_units_per_account": "Unità medie per host B2C - IPOTESI",
}
FC_LABELS = {
    "pre_live_platform": "Piattaforma prima del go-live (€/mese)",
    "live_platform": "Piattaforma in produzione, senza NextPax (€/mese)",
    "vercel_pro": "Vercel Pro (€/mese)",
    "company_admin": "Commercialista, PEC, banca (€/mese) - IPOTESI",
    "insurance": "Assicurazioni (€/mese) - IPOTESI",
    "company_from_month": "Mese da cui la società ha costi",
    "one_off_legal": "Pareri e contratti una tantum (€) - IPOTESI",
    "one_off_legal_month": "Mese della spesa una tantum",
    "marketing": "Marketing e vendite dal go-live (€/mese) - IPOTESI",
}
SC_LABELS = {
    "first_fp_start": "Primo Founding Partner: mese di avvio del pilota (99999 = mai)",
    "first_fp_units": "Primo Founding Partner: unità a regime",
    "first_fp_pilot_units": "Unità del pilota",
    "first_fp_pilot_months": "Durata del pilota (mesi)",
    "first_fp_ramp_months": "Durata della migrazione dopo il pilota (mesi)",
    "nextpax_min_from_full": "Minimo NextPax solo da migrazione completa (1 = sì)",
    "fp_start": "Primo mese altri Founding Partner",
    "fp_slots": "Posti Founding Partner oltre il primo",
    "fp_per_month": "Nuovi Founding Partner al mese",
    "list_start": "Primo mese agenzie a listino",
    "list_per_month": "Nuove agenzie a listino al mese",
    "units_per_agency": "Unità medie per agenzia",
    "agency_churn": "Churn mensile agenzie a listino",
    "b2c_start": "Primo mese B2C a pagamento",
    "b2c_per_month": "Nuovi host B2C al mese",
    "b2c_churn": "Churn mensile B2C",
    "founder_pay": "Compenso fondatore dall'avvio del pilota (€/mese, costo azienda)",
    "support_hire_units": "Assunzione supporto oltre N unità B2B",
    "support_cost": "Costo supporto (€/mese)",
}


def export(path: str, scenarios: list[Scenario], ue: UnitEconomics, fc: FixedCosts) -> None:
    wb = Workbook()
    ip = wb.active
    ip.title = "Ipotesi"
    ip["A1"] = "CasaZen - modello economico 36 mesi (nov 2026 - ott 2029). Celle gialle = ipotesi modificabili."
    ip["A1"].font = HEAD
    ip["A2"] = "Fonti: PR #526 §5-§6 per prezzi e costi unitari; il resto è IPOTESI (vedi business-plan.md §10)."
    ip["B4"] = "Scenario attivo (1-3)"
    ip["C4"] = 3
    ip["C4"].fill = INPUT
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(len(scenarios)))
    ip.add_data_validation(dv)
    dv.add("C4")

    names: dict[str, str] = {}
    r = 6
    ip.cell(r, 1, "Prezzi e costi unitari").font = HEAD
    r += 1
    for f in fields(UnitEconomics):
        ip.cell(r, 2, UE_LABELS[f.name])
        c = ip.cell(r, 3, getattr(ue, f.name))
        c.fill = INPUT
        names[f.name] = f"Ipotesi!$C${r}"
        r += 1
    r += 1
    ip.cell(r, 1, "Costi fissi").font = HEAD
    r += 1
    for f in fields(FixedCosts):
        ip.cell(r, 2, FC_LABELS[f.name])
        c = ip.cell(r, 3, getattr(fc, f.name))
        c.fill = INPUT
        names[f.name] = f"Ipotesi!$C${r}"
        r += 1
    r += 1
    ip.cell(r, 1, "Scenari").font = HEAD
    for i, s in enumerate(scenarios):
        ip.cell(r, 3 + i, s.name).font = HEAD
    ip.cell(r, 3 + len(scenarios), "Attivo").font = HEAD
    r += 1
    first_sc_col = get_column_letter(3)
    last_sc_col = get_column_letter(2 + len(scenarios))
    active_col = get_column_letter(3 + len(scenarios))
    for f in fields(Scenario):
        if f.name == "name":
            continue
        ip.cell(r, 2, SC_LABELS[f.name])
        for i, s in enumerate(scenarios):
            v = getattr(s, f.name)
            c = ip.cell(r, 3 + i, 99999 if v is None else int(v) if isinstance(v, bool) else v)  # None = mai
            c.fill = INPUT
        ip[f"{active_col}{r}"] = f"=INDEX({first_sc_col}{r}:{last_sc_col}{r},1,$C$4)"
        names[f.name] = f"Ipotesi!${active_col}${r}"
        r += 1
    ip.column_dimensions["A"].width = 4
    ip.column_dimensions["B"].width = 58
    for col in range(3, 4 + len(scenarios)):
        ip.column_dimensions[get_column_letter(col)].width = 16

    n = names
    ms = wb.create_sheet("Mensile")
    ms["A1"] = "Mese"
    ms["A1"].font = HEAD
    rows = [
        ("m", "Indice mese"),
        ("live", "Pilota del primo FP avviato (1/0)"),
        ("full", "Primo FP a regime (1/0)"),
        ("q_units", "Unità del primo Founding Partner"),
        ("fp", "Altri Founding Partner (agenzie)"),
        ("lst", "Agenzie a listino"),
        ("b2c", "Host B2C paganti"),
        ("fp_units", "Unità altri FP"),
        ("list_units", "Unità a listino"),
        ("b2b_units", "Unità B2B totali"),
        ("list_price", "Prezzo listino applicato (€/unità)"),
        ("rev_fee", "Ricavi canoni B2B"),
        ("rev_verif", "Ricavi verifiche al costo (FP)"),
        ("rev_b2c", "Ricavi B2C"),
        ("revenue", "RICAVI"),
        ("c_nextpax", "NextPax"),
        ("c_services", "SDI, AI, video"),
        ("c_verif", "Verifiche d'identità"),
        ("c_b2c", "Costi variabili B2C"),
        ("gm", "MARGINE LORDO"),
        ("c_fixed", "Costi fissi"),
        ("c_people", "Persone"),
        ("ebitda", "EBITDA"),
        ("timing", "Anticipo annuale FP (cassa)"),
        ("cash", "Cassa cumulata"),
    ]
    row_of = {key: i + 2 for i, (key, _) in enumerate(rows)}
    for key, label in rows:
        cell = ms.cell(row_of[key], 1, label)
        if label.isupper():
            cell.font = HEAD
    ms.column_dimensions["A"].width = 36

    for m in range(1, MONTHS + 1):
        col = get_column_letter(m + 1)
        prev = get_column_letter(m)
        ms[f"{col}1"] = month_label(m)
        ms[f"{col}1"].font = HEAD
        ms[f"{col}1"].alignment = Alignment(horizontal="center")
        R = {k: f"{col}{v}" for k, v in row_of.items()}
        P = {k: f"{prev}{v}" for k, v in row_of.items()}
        f = {}
        f["m"] = m
        start, pm, rm = n["first_fp_start"], n["first_fp_pilot_months"], n["first_fp_ramp_months"]
        qu, pu = n["first_fp_units"], n["first_fp_pilot_units"]
        full_m = f"({start}+{pm}+{rm}-1)"
        f["live"] = f"=IF({R['m']}>={start},1,0)"
        f["full"] = f"=IF({R['m']}>={full_m},1,0)"
        f["q_units"] = (f"=IF({R['m']}<{start},0,IF({R['m']}<{start}+{pm},{pu},"
                        f"MIN({qu},{pu}+({qu}-{pu})*({R['m']}-({start}+{pm})+1)/{rm})))")
        fp_prev = "0" if m == 1 else P["fp"]
        lst_prev = "0" if m == 1 else P["lst"]
        b2c_prev = "0" if m == 1 else P["b2c"]
        f["fp"] = (f"=IF(AND({R['m']}>={n['fp_start']},{R['live']}=1),"
                   f"MIN({n['fp_slots']},{fp_prev}+{n['fp_per_month']}),{fp_prev})")
        f["lst"] = (f"=IF(AND({R['m']}>={n['list_start']},{R['live']}=1),"
                    f"{lst_prev}*(1-{n['agency_churn']})+{n['list_per_month']},{lst_prev})")
        f["b2c"] = (f"=IF({R['m']}>={n['b2c_start']},"
                    f"{b2c_prev}*(1-{n['b2c_churn']})+{n['b2c_per_month']},{b2c_prev})")
        f["fp_units"] = f"={R['fp']}*{n['units_per_agency']}"
        f["list_units"] = f"={R['lst']}*{n['units_per_agency']}"
        f["b2b_units"] = f"={R['q_units']}+{R['fp_units']}+{R['list_units']}"
        f["list_price"] = (f"=IF({n['units_per_agency']}<50,{n['list_all_in_small']},"
                           f"IF({n['units_per_agency']}<150,{n['list_all_in_mid']},{n['list_all_in_large']}))")
        f["rev_fee"] = f"=({R['q_units']}+{R['fp_units']})*{n['fp_fee_per_unit']}+{R['list_units']}*{R['list_price']}"
        f["rev_verif"] = f"=({R['q_units']}+{R['fp_units']})*{n['verifications_per_unit']}*{n['verification_cost']}"
        f["rev_b2c"] = f"={R['b2c']}*{n['b2c_arpu']}"
        f["revenue"] = f"={R['rev_fee']}+{R['rev_verif']}+{R['rev_b2c']}"
        cm = f"({R['b2b_units']}+{R['b2c']}*{n['b2c_units_per_account']})"
        f["c_nextpax"] = (f"=IF({cm}>0,IF(AND({n['nextpax_min_from_full']}=1,{R['full']}=0),{cm}*{n['nextpax_per_unit']},"
                          f"{n['nextpax_min']}+MAX(0,{cm}-{n['nextpax_included_units']})*{n['nextpax_per_unit']}),0)")
        f["c_services"] = f"={R['b2b_units']}*{n['services_per_unit']}"
        f["c_verif"] = f"={R['b2b_units']}*{n['verifications_per_unit']}*{n['verification_cost']}"
        f["c_b2c"] = f"={R['b2c']}*{n['b2c_units_per_account']}*{n['services_per_unit']}"
        f["gm"] = f"={R['revenue']}-{R['c_nextpax']}-{R['c_services']}-{R['c_verif']}-{R['c_b2c']}"
        f["c_fixed"] = (f"=IF({R['live']}=1,{n['live_platform']}+{n['vercel_pro']}+{n['marketing']},{n['pre_live_platform']})"
                        f"+IF({R['m']}>={n['company_from_month']},{n['company_admin']}+{n['insurance']},0)"
                        f"+IF({R['m']}={n['one_off_legal_month']},{n['one_off_legal']},0)")
        f["c_people"] = (f"={R['live']}*{n['founder_pay']}+IF(AND({n['support_hire_units']}>0,"
                         f"{R['b2b_units']}>={n['support_hire_units']}),{n['support_cost']},0)")
        f["ebitda"] = f"={R['gm']}-{R['c_fixed']}-{R['c_people']}"
        f["timing"] = (f"=IF(AND({R['full']}=1,MOD({R['m']}-{full_m},12)=0),"
                       f"{qu}*{n['fp_fee_per_unit']}*12,0)-{R['full']}*{qu}*{n['fp_fee_per_unit']}")
        cash_prev = "0" if m == 1 else P["cash"]
        f["cash"] = f"={cash_prev}+{R['ebitda']}+{R['timing']}"
        for key, formula in f.items():
            c = ms[R[key]]
            c.value = formula
            if key not in ("m", "live", "full", "fp", "lst", "b2c"):
                c.number_format = '#,##0;[Red]-#,##0'
        ms.column_dimensions[col].width = 11
    ms.freeze_panes = "B2"

    sy = wb.create_sheet("Sintesi")
    sy["A1"] = "Sintesi annuale dello scenario attivo"
    sy["A1"].font = HEAD
    sy["B2"], sy["C2"], sy["D2"] = "Anno 1 (nov 26-ott 27)", "Anno 2", "Anno 3"
    for c in ("B2", "C2", "D2"):
        sy[c].font = HEAD
    summary = [
        ("Unità B2B a fine anno", "b2b_units", "last"),
        ("Host B2C a fine anno", "b2c", "last"),
        ("Ricavi", "revenue", "sum"),
        ("Margine lordo", "gm", "sum"),
        ("Costi fissi", "c_fixed", "sum"),
        ("Persone", "c_people", "sum"),
        ("EBITDA", "ebitda", "sum"),
        ("Ricavi mensili a fine anno", "revenue", "last"),
        ("Cassa cumulata a fine anno", "cash", "last"),
    ]
    for i, (label, key, how) in enumerate(summary):
        rr = i + 3
        sy.cell(rr, 1, label)
        for y in range(3):
            a = get_column_letter(2 + y * 12)
            b = get_column_letter(13 + y * 12)
            ref = f"Mensile!{a}{row_of[key]}:{b}{row_of[key]}" if how == "sum" else f"Mensile!{b}{row_of[key]}"
            c = sy.cell(rr, 2 + y, f"=SUM({ref})" if how == "sum" else f"={ref}")
            c.number_format = '#,##0;[Red]-#,##0'
    sy.cell(len(summary) + 4, 1, "Cassa cumulata minima")
    lo = sy.cell(len(summary) + 4, 2, f"=MIN(0,MIN(Mensile!B{row_of['cash']}:{get_column_letter(MONTHS + 1)}{row_of['cash']}))")
    lo.number_format = '#,##0;[Red]-#,##0'
    sy.column_dimensions["A"].width = 32
    for col in "BCD":
        sy.column_dimensions[col].width = 22
    wb.save(path)
