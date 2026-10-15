"""Ricalcola l'Excel con LibreOffice per ogni scenario e lo confronta con il modello Python."""

import os
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

from export_xlsx import export
from modello_economico import SCENARIOS, FixedCosts, UnitEconomics, run, yearly


def main(path: str) -> int:
    export(path, SCENARIOS, UnitEconomics(), FixedCosts())
    errors = 0
    with tempfile.TemporaryDirectory() as tmp:
        for idx, s in enumerate(SCENARIOS, start=1):
            wb = load_workbook(path)
            wb["Ipotesi"]["C4"] = idx
            src = os.path.join(tmp, f"s{idx}.xlsx")
            wb.save(src)
            out = os.path.join(tmp, "out")
            subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, src],
                           check=True, capture_output=True)
            vals = load_workbook(os.path.join(out, f"s{idx}.xlsx"), data_only=True)["Sintesi"]
            py = yearly(run(s))
            for y in range(3):
                for row, key in ((5, "revenue"), (6, "gross_margin"), (9, "ebitda"), (3, "units_end")):
                    xl = vals.cell(row, 2 + y).value
                    if xl is None or abs(xl - py[y][key]) > 0.5:
                        print(f"DIFF {s.name} anno {y + 1} {key}: xlsx={xl} python={py[y][key]:.2f}")
                        errors += 1
            print(f"{s.name}: confrontato")
    print("OK" if errors == 0 else f"{errors} differenze")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
